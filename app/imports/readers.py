"""Read-only Excel/CSV access: list sheets, read headers and rows.

Source files are opened read-only (``read_only=True`` for Excel; a plain text
open for CSV) and never written to - the application produces new files for
exports, never mutates the file a coordinator handed it (spec rule 55).
"""

from __future__ import annotations

import csv
from dataclasses import dataclass
from datetime import date, datetime
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

EXCEL_SUFFIXES = {".xlsx", ".xlsm"}
CSV_SUFFIXES = {".csv"}
SUPPORTED_SUFFIXES = EXCEL_SUFFIXES | CSV_SUFFIXES


class UnsupportedFileError(ValueError):
    """Raised when a file extension isn't one the import engine reads."""


def list_sheet_names(path: Path) -> list[str]:
    """List worksheet names; a CSV file has exactly one implicit "sheet"."""
    suffix = path.suffix.lower()
    if suffix in CSV_SUFFIXES:
        return [path.stem]
    if suffix in EXCEL_SUFFIXES:
        workbook = load_workbook(path, read_only=True, data_only=True)
        try:
            return list(workbook.sheetnames)
        finally:
            workbook.close()
    raise UnsupportedFileError(f"Unsupported file type: {suffix}")


@dataclass(slots=True)
class SheetData:
    """Headers plus row dicts read from one sheet/CSV file."""

    headers: list[str]
    rows: list[dict[str, Any]]


def _normalize_cell(value: Any) -> Any:
    """Normalize a raw cell value: strip strings, leave dates/numbers as-is."""
    if isinstance(value, str):
        stripped = value.strip()
        return stripped if stripped else None
    return value


def read_sheet(path: Path, sheet_name: str, *, max_rows: int | None = None) -> SheetData:
    """Read one sheet (or the whole file, for CSV) as a header row + row dicts.

    The first non-empty row is treated as the header row. Completely blank
    rows are skipped rather than becoming phantom "empty" records.
    """
    suffix = path.suffix.lower()
    if suffix in CSV_SUFFIXES:
        return _read_csv(path, max_rows=max_rows)
    if suffix in EXCEL_SUFFIXES:
        return _read_excel_sheet(path, sheet_name, max_rows=max_rows)
    raise UnsupportedFileError(f"Unsupported file type: {suffix}")


def _read_csv(path: Path, *, max_rows: int | None) -> SheetData:
    with path.open(newline="", encoding="utf-8-sig") as handle:
        reader = csv.reader(handle)
        rows_raw = list(reader)
    return _rows_to_sheet_data(rows_raw, max_rows=max_rows)


def _read_excel_sheet(path: Path, sheet_name: str, *, max_rows: int | None) -> SheetData:
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        worksheet = workbook[sheet_name]
        rows_raw = [list(row) for row in worksheet.iter_rows(values_only=True)]
    finally:
        workbook.close()
    return _rows_to_sheet_data(rows_raw, max_rows=max_rows)


def _rows_to_sheet_data(rows_raw: list[list[Any]], *, max_rows: int | None) -> SheetData:
    non_empty = [row for row in rows_raw if any(_normalize_cell(c) is not None for c in row)]
    if not non_empty:
        return SheetData(headers=[], rows=[])

    header_row = non_empty[0]
    headers = [str(_normalize_cell(h) or f"Column {i + 1}") for i, h in enumerate(header_row)]

    data_rows = non_empty[1:]
    if max_rows is not None:
        data_rows = data_rows[:max_rows]

    rows: list[dict[str, Any]] = []
    for raw_row in data_rows:
        row_dict: dict[str, Any] = {}
        for index, header in enumerate(headers):
            value = raw_row[index] if index < len(raw_row) else None
            row_dict[header] = _normalize_cell(value)
        rows.append(row_dict)
    return SheetData(headers=headers, rows=rows)


def coerce_date(value: Any) -> date | None:
    """Best-effort coercion of a cell value to a :class:`date`.

    Excel/openpyxl already hands back ``datetime``/``date`` objects for
    date-formatted cells; this also accepts common string formats for CSV
    input and manually-typed values.
    """
    if value is None or value == "":
        return None
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    text = str(value).strip()
    for fmt in ("%m/%d/%Y", "%Y-%m-%d", "%m-%d-%Y", "%m/%d/%y"):
        try:
            return datetime.strptime(text, fmt).date()
        except ValueError:
            continue
    return None


def coerce_float(value: Any) -> float | None:
    """Best-effort coercion of a cell value to a float."""
    if value is None or value == "":
        return None
    if isinstance(value, (int, float)):
        return float(value)
    text = str(value).strip().replace(",", "")
    try:
        return float(text)
    except ValueError:
        return None
