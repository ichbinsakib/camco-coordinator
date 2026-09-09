"""Export a :class:`Report` to an .xlsx workbook, one sheet per section."""

from __future__ import annotations

from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from app.reports.model import Report

_HEADER_FILL = PatternFill(start_color="1C2733", end_color="1C2733", fill_type="solid")
_HEADER_FONT = Font(color="FFFFFF", bold=True)
_TITLE_FONT = Font(bold=True, size=14)
_SUBTITLE_FONT = Font(italic=True, color="5B6675")


def _safe_sheet_title(title: str, used: set[str]) -> str:
    """Excel sheet names are limited to 31 chars and must be unique per workbook."""
    base = "".join(c for c in title if c not in r'[]:*?/\\')[:31]
    candidate = base or "Sheet"
    suffix = 1
    while candidate in used:
        suffix += 1
        candidate = f"{base[:28]}_{suffix}"
    used.add(candidate)
    return candidate


def _write_section(ws: Worksheet, report: Report, section) -> None:
    ws["A1"] = report.title
    ws["A1"].font = _TITLE_FONT
    ws["A2"] = report.subtitle or f"Generated {report.generated_at:%m/%d/%Y %I:%M %p} by {report.generated_by}"
    ws["A2"].font = _SUBTITLE_FONT
    ws["A3"] = section.title
    ws["A3"].font = Font(bold=True, size=12)

    header_row = 5
    for col_index, header in enumerate(section.headers, start=1):
        cell = ws.cell(row=header_row, column=col_index, value=header)
        cell.fill = _HEADER_FILL
        cell.font = _HEADER_FONT

    for row_offset, row_values in enumerate(section.rows, start=1):
        for col_index, value in enumerate(row_values, start=1):
            ws.cell(row=header_row + row_offset, column=col_index, value=value)

    if not section.rows:
        ws.cell(row=header_row + 1, column=1, value="(no records)")

    for col_index in range(1, len(section.headers) + 1):
        column_letter = get_column_letter(col_index)
        width = max(
            (len(str(section.headers[col_index - 1])), *(len(str(r[col_index - 1])) for r in section.rows if col_index - 1 < len(r))),
            default=10,
        )
        ws.column_dimensions[column_letter].width = min(max(width + 2, 10), 50)

    ws.freeze_panes = ws.cell(row=header_row + 1, column=1)
    ws["A2"].alignment = Alignment(wrap_text=False)


def export_report_to_excel(report: Report, output_path: Path) -> Path:
    """Write ``report`` to an .xlsx file at ``output_path``, one sheet per section."""
    workbook = Workbook()
    workbook.remove(workbook.active)

    used_titles: set[str] = set()
    for section in report.sections:
        sheet_name = _safe_sheet_title(section.title, used_titles)
        ws = workbook.create_sheet(sheet_name)
        _write_section(ws, report, section)

    if not report.sections:
        workbook.create_sheet("Report")

    output_path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(output_path)
    return output_path
