"""Integration tests for report builders and exporters."""

from __future__ import annotations

from datetime import UTC, date, timedelta
from pathlib import Path

from sqlalchemy.orm import Session

from app.config.constants import OrderStatus
from app.config.settings import AppSettings
from app.models.core import Customer, Part
from app.models.orders import CustomerOrder, CustomerOrderLine
from app.reports.builders import (
    build_customer_order_report,
    build_daily_coordinator_report,
    build_management_summary_report,
)
from app.reports.csv_export import export_report_to_csv
from app.reports.excel_export import export_report_to_excel
from app.reports.model import Report, ReportSection
from app.reports.pdf_export import export_report_to_pdf


def _seed_past_due_line(session: Session) -> None:
    customer = Customer(code="ACME", name="Acme Corp", importance=1)
    part = Part(part_number="P1", revision="A")
    session.add_all([customer, part])
    session.flush()
    order = CustomerOrder(co_number="CO-1", customer_id=customer.id)
    session.add(order)
    session.flush()
    session.add(
        CustomerOrderLine(
            customer_order_id=order.id, line_number=1, part_id=part.id, quantity_ordered=10,
            due_date=date.today() - timedelta(days=20), status=OrderStatus.IN_PROGRESS.value,
        )
    )
    session.commit()


def test_daily_coordinator_report_includes_past_due_line(db_session: Session) -> None:
    _seed_past_due_line(db_session)
    report = build_daily_coordinator_report(db_session, AppSettings(), generated_by="Tester")
    past_due_section = next(s for s in report.sections if s.title == "Past Due Orders")
    assert len(past_due_section.rows) == 1
    assert past_due_section.rows[0][0] == "CO-1"


def test_daily_coordinator_report_critical_section_escalates(db_session: Session) -> None:
    _seed_past_due_line(db_session)
    report = build_daily_coordinator_report(db_session, AppSettings(), generated_by="Tester")
    critical_section = next(s for s in report.sections if s.title == "Critical Priority Issues")
    assert len(critical_section.rows) == 1  # 20 days late, importance=1 -> CRITICAL


def test_customer_order_report_empty_when_no_orders(db_session: Session) -> None:
    report = build_customer_order_report(db_session, AppSettings(), generated_by="Tester")
    assert report.sections[0].rows == []


def test_management_summary_includes_kpis(db_session: Session) -> None:
    _seed_past_due_line(db_session)
    report = build_management_summary_report(db_session, AppSettings(), generated_by="Tester")
    kpi_section = report.sections[0]
    kpi_dict = dict(kpi_section.rows)
    assert kpi_dict["Past Due"] == "1"


def _sample_report() -> Report:
    from datetime import datetime

    return Report(
        title="Sample Report",
        generated_at=datetime.now(UTC),
        generated_by="Tester",
        sections=[
            ReportSection("Section One", ["A", "B"], [["1", "2"], ["3", "4"]]),
            ReportSection("Section Two", ["X"], []),
        ],
    )


def test_excel_export_creates_one_sheet_per_section(tmp_path: Path) -> None:
    report = _sample_report()
    path = export_report_to_excel(report, tmp_path / "out.xlsx")
    assert path.exists()

    from openpyxl import load_workbook

    wb = load_workbook(path)
    assert wb.sheetnames == ["Section One", "Section Two"]
    assert wb["Section One"]["A5"].value == "A"


def test_pdf_export_creates_file(tmp_path: Path) -> None:
    report = _sample_report()
    path = export_report_to_pdf(report, tmp_path / "out.pdf")
    assert path.exists()
    assert path.stat().st_size > 0


def test_csv_export_creates_one_file_per_section(tmp_path: Path) -> None:
    report = _sample_report()
    paths = export_report_to_csv(report, tmp_path / "out.csv")
    assert len(paths) == 2
    assert all(p.exists() for p in paths)
    content = paths[0].read_text(encoding="utf-8-sig")
    assert "Section One" in content or "A,B" in content


def test_csv_export_single_section_uses_base_name(tmp_path: Path) -> None:
    from datetime import datetime

    report = Report(
        title="Single",
        generated_at=datetime.now(UTC),
        generated_by="Tester",
        sections=[ReportSection("Only", ["A"], [["1"]])],
    )
    paths = export_report_to_csv(report, tmp_path / "single.csv")
    assert len(paths) == 1
    assert paths[0] == tmp_path / "single.csv"
